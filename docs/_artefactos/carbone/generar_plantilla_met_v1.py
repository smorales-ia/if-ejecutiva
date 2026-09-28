#!/usr/bin/env python3
"""
Genera docs/_artefactos/carbone/PLANTILLA_MET_v1.docx — plantilla Carbone del informe
MetLife (gold master MET-6283), según el PLANO del Agente 2 y el ORÁCULO del Agente 4
(T-PDF-IMPRENTA-20260927 · docs/_evidencia/T-PDF-IMPRENTA-20260927/{plano-plantilla,oraculo-met6283}.md).

Reglas duras implementadas:
- Cada tag Carbone {d.*} va en UN run único (jamás partido entre runs).
- Los 8 grupos {d.cualitativa.*} NO se insertan (son Record|null → imprimirían [object Object]);
  sus zonas quedan como layout fijo con celdas vacías (P1-1/P1-2).
- Placeholders de imagen: PNG gris con alt-text (docPr/@descr) = tag Carbone de imagen.
- Colores del gold master: azul #095085 · celeste #8DB4E2 · gris #D4D4D4 (medidos en el XLSM).
- verificar(): los 52 tags de variables_requeridas.txt (tanda BASE) presentes; exit≠0 si falta alguno.
"""
import re
import struct
import sys
import zipfile
import zlib
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AQUI = Path(__file__).resolve().parent
REPO = AQUI.parents[2]
SALIDA = AQUI / "PLANTILLA_MET_v1.docx"
REQUERIDAS = REPO / "docs/_evidencia/T-PDF-BASE-20260927/variables_requeridas.txt"
PDF_REF = REPO / "docs/_referencias/Informe FRANCISCO VERGARA UNDURRAGA_Met6283.pdf"

AZUL = "095085"
CELESTE = "8DB4E2"
GRIS = "D4D4D4"
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
NEGRO = RGBColor(0, 0, 0)

# Los 8 grupos cualitativa.* que NO se insertan (lista blanca de verificar()).
CUALITATIVA_EXCLUIDOS = {
    "{d.cualitativa.normativa}", "{d.cualitativa.sector}", "{d.cualitativa.geometriaTerreno}",
    "{d.cualitativa.emplazamiento}", "{d.cualitativa.constructivas}", "{d.cualitativa.servicios}",
    "{d.cualitativa.comodidades}", "{d.cualitativa.detallePropiedad}",
}

DECLARACION = (
    "El profesional que firma declara que no tiene hoy, ni espera tener en el futuro, interés en la "
    "propiedad tasada ni ningún impedimento para llevar a cabo este trabajo. No tiene personal interés "
    "ni participación en los usos que se hagan de la tasación ni con las personas que participen en la "
    "operación. Ha inspeccionado la vivienda y la información que en esta fecha presenta es totalmente "
    "verdadera, y no ha olvidado nada de importancia. Los inconvenientes y limitaciones que pueda tener "
    "la vivienda y su vecindario están mencionados. Por otra parte se mantendrá un nivel de "
    "confidencialidad de la información obtenida, acorde a las exigencias de la Institución, que desde "
    "ya declara conocer y aceptar."
)
BOILERPLATE_REFERENCIAS = (
    "Las muestras fueron seleccionadas por ser propiedades que comparten características similares en "
    "cuanto a dimensiones, programa, tipología de edificación, sector y se homologa en función a sus "
    "características propias como; diseño constructivo, calidad de sus materiales, calidad de sus "
    "revestimientos interiores y estado de mantención."
)

# ----------------------------------------------------------------------------- helpers OOXML


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def vertical(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    td = OxmlElement("w:textDirection")
    td.set(qn("w:val"), "btLr")
    tcPr.append(td)


def bordes(table, sz=4):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)


def sin_bordes(table):
    tbl = table._tbl
    borders = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:val"), "none")
        borders.append(el)
    tbl.tblPr.append(borders)


def run_fmt(run, size, bold, color, font="Calibri"):
    run.font.name = font
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def celda(cell, texto, size=6, bold=False, color=None, align=None, font="Calibri"):
    """Escribe la celda con UN run único (regla de tags Carbone)."""
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    run = p.add_run(texto)
    run_fmt(run, size, bold, color, font)
    if align:
        p.alignment = align
    return p


def celda_partes(cell, partes, align=None):
    """partes = [(texto, size, bold, color)] — cada tag completo en su propio run."""
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    for texto, size, bold, color in partes:
        run = p.add_run(texto)
        run_fmt(run, size, bold, color)
    if align:
        p.alignment = align
    return p


def parrafo(doc, texto, size=6, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, color=None,
            space_after=2):
    p = doc.add_paragraph()
    run = p.add_run(texto)
    run_fmt(run, size, bold, color)
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    return p


def barra(doc, texto, size=10, fill=AZUL, color=BLANCO):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    t.columns[0].width = Cm(19.5)
    shade(t.cell(0, 0), fill)
    celda(t.cell(0, 0), texto, size=size, bold=True, color=color,
          align=WD_ALIGN_PARAGRAPH.CENTER)
    return t


PNG_GRIS = None


def png_gris(path):
    """PNG gris 120×80 sin dependencias (zlib + struct)."""
    w, h = 120, 80
    fila = b"\x00" + b"\xd9\xd9\xd9" * w
    raw = fila * h

    def chunk(tipo, data):
        c = struct.pack(">I", len(data)) + tipo + data
        return c + struct.pack(">I", zlib.crc32(tipo + data) & 0xFFFFFFFF)

    png = (b"\x89PNG\r\n\x1a\n"
           + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
           + chunk(b"IDAT", zlib.compress(raw))
           + chunk(b"IEND", b""))
    path.write_bytes(png)
    return path


def imagen(cell_o_doc, alt, ancho_cm, alto_cm=None, path=None):
    """Inserta imagen (placeholder gris por defecto) con alt-text = tag Carbone."""
    src = path or PNG_GRIS
    if hasattr(cell_o_doc, "paragraphs"):  # celda
        p = cell_o_doc.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
    else:
        p = cell_o_doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
    shape = run.add_picture(str(src), width=Cm(ancho_cm),
                            height=Cm(alto_cm) if alto_cm else None)
    if alt:
        shape._inline.docPr.set("descr", alt)
    return shape


def logo_portada():
    """Extrae el logo embebido de la página 1 del PDF de referencia."""
    try:
        import pymupdf
        pdf = pymupdf.open(str(PDF_REF))
        imgs = pdf[0].get_images(full=True)
        if not imgs:
            return None
        d = pdf.extract_image(imgs[0][0])
        out = AQUI / f"_logo_extraido.{d['ext']}"
        out.write_bytes(d["image"])
        return out
    except Exception as e:  # noqa: BLE001
        print(f"  WARN logo no extraíble ({e}); uso placeholder", file=sys.stderr)
        return None


def encabezado_hoja(doc, n):
    """Encabezado repetido de las Hojas 2-7 (primera tabla de cada hoja) + folio."""
    t = doc.add_table(rows=2, cols=5)
    t.autofit = False
    bordes(t)
    anchos = (3.2, 6.3, 2.6, 5.4, 2.0)
    for fila in t.rows:
        for i, c in enumerate(fila.cells):
            c.width = Cm(anchos[i])
    celda(t.cell(0, 0), "Nombre Cliente", 7, True, None)
    celda(t.cell(0, 1), "{d.partes.propietario}", 7, True)
    celda(t.cell(0, 2), "Dirección", 7, True)
    celda(t.cell(0, 3), "{d.propiedad.direccion}", 7, True)
    celda(t.cell(0, 4), f"Hoja N°{n}", 6, True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    celda(t.cell(1, 0), "RUT", 7, True)
    celda(t.cell(1, 1), "{d.partes.rut}", 7, True)
    celda(t.cell(1, 2), "Nº Interno", 7, True)
    celda(t.cell(1, 3), "{d.meta.numeroSolicitudCliente}", 7, True)
    celda(t.cell(1, 4), "", 6)
    for i, c in enumerate(t.columns):
        pass
    return t


def bloque_con_sidebar(doc, rotulo, filas, cols, anchos_cm):
    """Tabla externa 1×2: sidebar azul vertical + tabla interna de contenido."""
    ext = doc.add_table(rows=1, cols=2)
    ext.autofit = False
    sin_bordes(ext)
    sb = ext.cell(0, 0)
    sb.width = Cm(0.55)
    shade(sb, AZUL)
    vertical(sb)
    celda(sb, rotulo, 6, True, BLANCO, align=WD_ALIGN_PARAGRAPH.CENTER)
    cont = ext.cell(0, 1)
    cont.width = Cm(18.9)
    inner = cont.add_table(rows=filas, cols=cols)
    inner.autofit = False
    bordes(inner)
    if anchos_cm:
        for fila in inner.rows:
            for i, c in enumerate(fila.cells):
                c.width = Cm(anchos_cm[i])
    return inner


# ----------------------------------------------------------------------------- páginas


def portada(doc):
    barra(doc, "INFORME DE TASACION", size=18)
    doc.add_paragraph()
    # recuadro logo
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    t.columns[0].width = Cm(13.5)
    bordes(t, sz=8)
    logo = logo_portada()
    if logo:
        imagen(t.cell(0, 0), None, 9.9, 5.8, path=logo)
    else:
        imagen(t.cell(0, 0), None, 9.9, 5.8)
        celda(t.cell(0, 0), "VALUE PROPERTY · Tasaciones Bienes Raíces", 11, True,
              align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()
    barra(doc, "ANTECEDENTES", size=14)
    doc.add_paragraph()
    ficha = doc.add_table(rows=7, cols=2)
    ficha.alignment = WD_TABLE_ALIGNMENT.CENTER
    ficha.autofit = False
    bordes(ficha, sz=8)
    datos = [
        ("Numero Solicitud :", "{d.meta.numeroSolicitudCliente}"),
        ("Institución :", "{d.clienteInforme.nombre}"),
        ("Nombre Cliente :", "{d.partes.propietario}"),
        ("Rut :", "{d.partes.rut}"),
        ("Dirección Propiedad :", "{d.propiedad.direccion}"),
        ("Comuna :", "{d.propiedad.comuna}"),
        ("Región  :", "{d.propiedad.region}"),
    ]
    for i, (k, v) in enumerate(datos):
        ficha.rows[i].cells[0].width = Cm(6.0)
        ficha.rows[i].cells[1].width = Cm(9.5)
        celda(ficha.cell(i, 0), k, 11, True)
        celda(ficha.cell(i, 1), v, 11, True)
    for _ in range(3):
        doc.add_paragraph()
    pie = doc.add_table(rows=2, cols=2)
    pie.alignment = WD_TABLE_ALIGNMENT.CENTER
    pie.autofit = False
    bordes(pie, sz=8)
    celda(pie.cell(0, 0), "www.valueproperty.cl", 11, True)
    celda(pie.cell(1, 0), "Mail: info@valueproperty.cl", 11, True)
    celda(pie.cell(0, 1), "Dirección: Santa Magdalena 75 Of 310, Providencia - Santiago", 11, True)
    celda(pie.cell(1, 1), "Fono: 22 500 0366", 11, True)


def hoja1(doc):
    barra(doc, "INFORME DE TASACION", size=10)
    # banda Nº interno + folio (incluye {d.meta.codigo} — cobertura de contrato)
    band = doc.add_table(rows=1, cols=2)
    band.autofit = False
    band.columns[0].width = Cm(15.5)
    band.columns[1].width = Cm(4.0)
    shade(band.cell(0, 0), AZUL)
    celda_partes(band.cell(0, 0), [
        ("Nº Interno  :  ", 7.5, True, BLANCO),
        ("{d.meta.numeroSolicitudCliente}", 7.5, True, BLANCO),
        ("   (", 7.5, False, BLANCO),
        ("{d.meta.codigo}", 7.5, False, BLANCO),
        (")", 7.5, False, BLANCO),
    ])
    celda(band.cell(0, 1), "Hoja N°1", 6, True, align=WD_ALIGN_PARAGRAPH.RIGHT)

    # ---- B1 Identificación (3 pares de columnas rótulo/valor)
    colA = [
        ("Cliente", "{d.partes.propietario}"), ("RUT Cliente", "{d.partes.rut}"),
        ("Propietario", "{d.partes.propietario}"), ("RUT Prop.", "{d.partes.rut}"),
        ("Ejecutivo", "{d.partes.ejecutivo}"), ("Tasador", "{d.partes.tasador.nombre}"),
        ("Objetivo", "{d.propiedad.objetivo}"), ("Tipo Prop.", "{d.propiedad.tipoPropiedad}"),
        ("Fecha Tasación", "{d.partes.fechaVisita:formatD(DD-MM-YYYY)}"),
        ("Destino SII", "{d.sii.destinoSii}"), ("Ejecutivo", ""), ("Formalizador", ""),
        ("N° SOLICITUD", "{d.meta.nOperacionCliente}"),
    ]
    colB = [
        ("Dirección", "{d.propiedad.direccion}"), ("Casa Nº", ""),
        ("Comuna", "{d.propiedad.comuna}"), ("Region", "{d.propiedad.region}"),
        ("Rol", "{d.sii.rolSii}"), ("Condominio", "{d.propiedad.proyectoCondominio}"),
        ("Pisos Propiedad", "{d.propiedad.pisos}"), ("Sitio/Lote", ""), ("Zona", ""),
        ("Manzana", "{d.sii.codManzana}"),
        ("Estacionamientos Asociados", "{d.propiedad.estacionamientos}"),
        ("Estacionamientos Asignados", ""), ("Bodegas", "{d.propiedad.bodegas}"),
    ]
    colC = [
        ("DFL-2", "{d.propiedad.dfl2}"), ("Año Construccion", "{d.propiedad.anioConstruccion}"),
        ("Vida util", "{d.propiedad.vidaUtil}"), ("Permiso", "{d.legales.permisoEdificacion}"),
        ("Recepcion", "{d.legales.recepcionFinal}"), ("Ampliación", "No"),
        ("Subterraneos", ""), ("Estado conservacion", "{d.propiedad.estadoConservacion:upperCase}"),
        ("Sello SEC", ""), ("Afecto a expropi.", ""), ("Fuente Información", ""),
        ("Mansarda", ""), ("Descripcion Expropiación", "{d.textosIA.textoExpropiacion}"),
    ]
    anchos = (2.1, 4.2, 2.3, 4.0, 2.3, 4.0)
    ident = bloque_con_sidebar(doc, "Identificación", len(colA) + 1, 6, anchos)
    for i in range(len(colA)):
        for j, col in enumerate((colA, colB, colC)):
            k, v = col[i]
            cl = ident.cell(i, j * 2)
            shade(cl, GRIS)
            celda(cl, k, 5.5, True)
            celda(ident.cell(i, j * 2 + 1), v, 5.5, False)
    # fila mini-mapa Ubicación/Empresa
    fila_mapa = ident.rows[len(colA)]
    celda(fila_mapa.cells[0], "Ubicación/Empresa", 5.5, True)
    shade(fila_mapa.cells[0], GRIS)
    imagen(fila_mapa.cells[1], "{d.mapa.staticMapUrl}", 3.2, 2.6)

    # ---- B2 Síntesis
    sint = bloque_con_sidebar(doc, "Síntesis de la Prop.", 2, 1, (18.9,))
    celda(sint.cell(0, 0), "{d.textosIA.sintesisPropiedad}", 5.5,
          align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    celda(sint.cell(1, 0), "{d.textosIA.descripcionSector}", 5.5,
          align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    # ---- B3/B4 Referencias
    def tabla_ref(titulo, col4, filtro):
        cab = ["N°", "Fecha", "Tipo", "Dirección referencias", "Año", col4, "Total UF",
               "Sup. Terreno.", "Sup. Constr.", "OO.CC.", "UF/m² T.", "UF/m² C.",
               "Comentarios Relevantes"]
        anchos_r = (0.6, 1.0, 0.9, 3.6, 0.8, 1.5, 1.3, 1.5, 1.4, 1.0, 1.0, 1.0, 3.3)
        t = bloque_con_sidebar(doc, "Referencias y Mercado", 6, len(cab), anchos_r)
        # título de bloque en la primera fila combinada
        fila_t = t.rows[0]
        fila_t.cells[0].merge(fila_t.cells[len(cab) - 1])
        shade(fila_t.cells[0], GRIS)
        celda(fila_t.cells[0], titulo, 6.5, True)
        for j, h in enumerate(cab):
            shade(t.cell(1, j), GRIS)
            celda(t.cell(1, j), h, 5.5, True, align=WD_ALIGN_PARAGRAPH.CENTER)
        # fila loop [i, filtro]
        pref = f"{{d.comparablesInforme.filas[i, tipoReferencia='{filtro}']"
        datos = ["", "", pref + ".tipoReferencia}", pref + ".direccion}", pref + ".anio}",
                 "", pref + ".precioUf:formatN(0)}", pref + ".supTerreno:formatN(2)}",
                 pref + ".supConstruida:formatN(2)}", "", "",
                 pref + ".ufM2Construccion:formatN(2)}", ""]
        for j, v in enumerate(datos):
            celda(t.cell(2, j), v, 5.5, align=WD_ALIGN_PARAGRAPH.RIGHT if j >= 6 else None)
        celda(t.cell(3, 0), "{d.comparablesInforme.filas[i+1].direccion}", 5.5)
        # agregados
        for j in range(len(cab)):
            shade(t.cell(4, j), CELESTE)
            shade(t.cell(5, j), CELESTE)
        celda(t.cell(4, 3), "PROMEDIO DE LA MUESTRA", 5.5, True)
        celda(t.cell(4, 11), "{d.comparablesInforme.promedioUfM2:formatN(2)}", 5.5, True,
              align=WD_ALIGN_PARAGRAPH.RIGHT)
        celda(t.cell(5, 3), "TASACION", 5.5, True)
        celda(t.cell(5, 6), "{d.terminales.valorComercialUf:formatN(0)}", 5.5, True,
              align=WD_ALIGN_PARAGRAPH.RIGHT)
        celda(t.cell(5, 7), "{d.propiedad.supTerrenoM2:formatN(0)}", 5.5, True,
              align=WD_ALIGN_PARAGRAPH.RIGHT)
        celda(t.cell(5, 8), "{d.propiedad.supConstruccionM2:formatN(0)}", 5.5, True,
              align=WD_ALIGN_PARAGRAPH.RIGHT)
        celda(t.cell(5, 11), "{d.comparablesInforme.tasacionUfM2:formatN(2)}", 5.5, True,
              align=WD_ALIGN_PARAGRAPH.RIGHT)
        return t

    tabla_ref("REF. OFERTAS", "Teléfono", "Oferta")
    # fila V/S en tabla propia (misma zona)
    vs = doc.add_table(rows=1, cols=2)
    vs.autofit = False
    bordes(vs)
    vs.columns[0].width = Cm(14.5)
    vs.columns[1].width = Cm(5.0)
    shade(vs.cell(0, 0), CELESTE)
    shade(vs.cell(0, 1), CELESTE)
    celda(vs.cell(0, 0), "TASACION V/S PROMEDIO DE LA MUESTRA", 5.5, True)
    celda_partes(vs.cell(0, 1), [("{d.comparablesInforme.tasacionVsPct:formatN(0)}", 5.5, True, None),
                                 ("%", 5.5, True, None)], align=WD_ALIGN_PARAGRAPH.RIGHT)

    tabla_ref("REF. C.B.R.", "Foja y Número", "CBR")

    # ---- ANALISIS DE RENTABILIDAD
    rent_datos = [
        ("Vida Util Remanente Años", "{d.propiedad.vidaUtil}"),
        ("Arriendo Bruto $ / mes", "{d.rentabilidad.arriendoBrutoMensualClp:formatN(0)}"),
        ("UF/ mes", "{d.rentabilidad.arriendoUfMes:formatN(1)}"),
        ("Gasto Anual $", "{d.rentabilidad.gastoAnualClp:formatN(0)}"),
        ("Tasa Exigida Proyecto", None),  # se arma con partes (:mul(100) + % fijo)
        ("Tiempo Renta (años)", ""),
        ("Ingreso Líquido Anual", "{d.rentabilidad.ingresoLiquidoAnualClp:formatN(0)}"),
        ("Renta Perpetua", "{d.rentabilidad.rentaPerpetuaClp:formatN(0)}"),
    ]
    rent = doc.add_table(rows=len(rent_datos) + 1, cols=2)
    rent.autofit = False
    bordes(rent)
    rent.columns[0].width = Cm(5.0)
    rent.columns[1].width = Cm(4.5)
    fila0 = rent.rows[0]
    fila0.cells[0].merge(fila0.cells[1])
    shade(fila0.cells[0], GRIS)
    celda(fila0.cells[0], "ANALISIS DE RENTABILIDAD", 6.5, True,
          align=WD_ALIGN_PARAGRAPH.CENTER)
    for i, (k, v) in enumerate(rent_datos, start=1):
        celda(rent.cell(i, 0), k, 5.5, True)
        if v is None:
            celda_partes(rent.cell(i, 1),
                         [("{d.rentabilidad.tasaCapRate:mul(100):formatN(1)}", 5.5, False, None),
                          ("%", 5.5, False, None)], align=WD_ALIGN_PARAGRAPH.RIGHT)
        else:
            celda(rent.cell(i, 1), v, 5.5, align=WD_ALIGN_PARAGRAPH.RIGHT)

    # ---- B5 boilerplate
    barra(doc, "ANÁLISIS DE LAS REFERENCIAS", size=6.5, fill=GRIS, color=NEGRO)
    parrafo(doc, BOILERPLATE_REFERENCIAS, size=5.5)

    # ---- B6 Cuadro de valoración
    cab_q = ["Item", "Nº", "Detalle Item valorado", "Rol SII", "Año", "Tipo",
             "Situación Municipal", "Estado", "Grntía.", "Origen Super.", "Superficie m²",
             "UF/m² Nuevo", "D. F.", "UF/m²", "Valor Comercial UF", "Valor Seguros UF",
             "Valor Liquidación UF"]
    anchos_q = (1.1, 0.5, 2.3, 1.1, 0.7, 0.6, 1.5, 0.7, 0.8, 1.2, 1.2, 1.0, 0.6, 0.9,
                1.5, 1.4, 1.5)
    q = bloque_con_sidebar(doc, "Valores y Firma", 10, len(cab_q), anchos_q)
    fila_tq = q.rows[0]
    fila_tq.cells[0].merge(fila_tq.cells[len(cab_q) - 1])
    shade(fila_tq.cells[0], GRIS)
    celda(fila_tq.cells[0], "CUADRO DE VALORACION", 6.5, True,
          align=WD_ALIGN_PARAGRAPH.CENTER)
    for j, h in enumerate(cab_q):
        shade(q.cell(1, j), GRIS)
        celda(q.cell(1, j), h, 5, True, align=WD_ALIGN_PARAGRAPH.CENTER)
    filaq = ["{d.cuadro.items[i].tipoItem}", "", "{d.cuadro.items[i].descripcion}",
             "{d.sii.rolSii}", "", "", "{d.cuadro.items[i].situacionMunicipal}", "",
             "{d.cuadro.items[i].aportaAGarantia:ifEQ(true):show(SI):elseShow( )}", "",
             "{d.cuadro.items[i].supM2:formatN(2)}",
             "{d.cuadro.items[i].ufM2Aplicado:formatN(2)}",
             "{d.cuadro.items[i].factorAplicado:formatN(2)}", "",
             "{d.cuadro.items[i].ufTotalItem:formatN(2)}", "", ""]
    for j, v in enumerate(filaq):
        celda(q.cell(2, j), v, 5, align=WD_ALIGN_PARAGRAPH.RIGHT if j >= 10 else None)
    celda(q.cell(3, 0), "{d.cuadro.items[i+1].tipoItem}", 5)
    for i, rot in enumerate(("TOTAL EDIFICACION", "TOTAL OBRAS COMPLEMENTARIAS",
                             "TOTAL TERRENO"), start=4):
        celda(q.cell(i, 2), rot, 5, True)
    fila_vc = 7
    shade(q.cell(fila_vc, 2), CELESTE)
    celda(q.cell(fila_vc, 2), "VALOR COMERCIAL NORMAL :", 5.5, True)
    celda(q.cell(fila_vc, 14), "{d.cuadro.totalUf:formatN(2)}", 5.5, True,
          align=WD_ALIGN_PARAGRAPH.RIGHT)
    celda(q.cell(fila_vc, 15), "{d.terminales.seguroIncendioUf:formatN(2)}", 5.5, True,
          align=WD_ALIGN_PARAGRAPH.RIGHT)
    celda(q.cell(fila_vc, 16), "{d.terminales.valorLiquidacionUf:formatN(2)}", 5.5, True,
          align=WD_ALIGN_PARAGRAPH.RIGHT)
    celda(q.cell(8, 2), "BIENES NO CONSIDERADOS GARANTIA :", 5, True)

    # ---- B8/B9/B10 Foto fachada + Valor tasación + firma
    dos = doc.add_table(rows=1, cols=2)
    dos.autofit = False
    sin_bordes(dos)
    dos.columns[0].width = Cm(9.2)
    dos.columns[1].width = Cm(10.3)
    izq = dos.cell(0, 0)
    der = dos.cell(0, 1)

    tfoto = izq.add_table(rows=2, cols=1)
    bordes(tfoto)
    shade(tfoto.cell(0, 0), GRIS)
    celda(tfoto.cell(0, 0), "FOTO FACHADA", 6.5, True, align=WD_ALIGN_PARAGRAPH.CENTER)
    imagen(tfoto.cell(1, 0), "{d.fotos.fotos[i=0].url}", 8.3, 4.2)

    firma_datos = [
        ("Tasador :", "{d.partes.tasador.nombre}"),
        ("Visador :", "{d.partes.visador.nombre}"),
        ("Fecha visita :", "{d.partes.fechaVisita:formatD(LL)}"),
        ("Fecha Visado :", "{d.partes.fechaVisado:formatD(LL)}"),
        ("REVISOR", "{d.clienteInforme.nombre}"),
        ("Fecha revisión :", ""),
    ]
    tfirma = izq.add_table(rows=len(firma_datos) + 1, cols=2)
    bordes(tfirma)
    for i, (k, v) in enumerate(firma_datos):
        celda(tfirma.cell(i, 0), k, 6, True)
        celda(tfirma.cell(i, 1), v, 6, True)
    celda(tfirma.cell(len(firma_datos), 0), "Firma :", 6, True)
    imagen(tfirma.cell(len(firma_datos), 1), "{d.partes.tasador.firmaUrl}", 2.6, 1.2)

    tv = der.add_table(rows=4, cols=1)
    bordes(tv)
    shade(tv.cell(0, 0), CELESTE)
    celda(tv.cell(0, 0), "VALOR TASACIÓN", 10, True, align=WD_ALIGN_PARAGRAPH.CENTER)
    shade(tv.cell(1, 0), CELESTE)
    celda_partes(tv.cell(1, 0), [
        ("UF  ", 8.5, True, None), ("{d.terminales.valorComercialUf:formatN(2)}", 8.5, True, None),
        ("      $  ", 8.5, True, None),
        ("{d.terminales.valorComercialClp:formatN(0)}", 8.5, True, None),
        ("      al ", 8.5, True, None),
        ("{d.partes.fechaVisita:formatD(DD-MM-YYYY)}", 8.5, True, None),
    ], align=WD_ALIGN_PARAGRAPH.CENTER)
    celda_partes(tv.cell(2, 0), [
        ("Velocidad de venta normal : ", 7, True, None),
        ("{d.propiedad.velocidadVentaEstimada:upperCase}", 7, True, None),
    ], align=WD_ALIGN_PARAGRAPH.CENTER)
    celda_partes(tv.cell(3, 0), [
        ("1UF =  $ ", 7, True, None), ("{d.terminales.ufDia:formatN(2)}", 7, True, None),
        ("        1US$ =  $ ", 7, True, None), ("{d.terminales.usdDia:formatN(2)}", 7, True, None),
    ], align=WD_ALIGN_PARAGRAPH.CENTER)

    val_filas = [
        ("Valor de Reposición :", "", "{d.terminales.valorReposicionUf:formatN(2)}",
         "{d.terminales.valorReposicionClp:formatN(0)}", None),
        ("Seguro Incendio y otros :", "", "{d.terminales.seguroIncendioUf:formatN(2)}",
         "{d.terminales.seguroIncendioClp:formatN(0)}", None),
        ("Avalúo fiscal propiedad :", "", "{d.terminales.avaluoFiscalUf:formatN(2)}",
         "{d.sii.avaluoTotal:formatN(0)}", None),
        ("Valor a Remate  65% :", "", "{d.terminales.valorRemateUf:formatN(2)}",
         "{d.terminales.valorRemateClp:formatN(0)}", CELESTE),
        ("Liquid. Normal  82,5% :", "", "{d.terminales.valorLiquidacionUf:formatN(2)}",
         "{d.terminales.valorLiquidacionClp:formatN(0)}", None),
    ]
    t5 = der.add_table(rows=len(val_filas) + 1, cols=4)
    bordes(t5)
    for j, h in enumerate(("", "US$", "UF", "$")):
        shade(t5.cell(0, j), GRIS)
        celda(t5.cell(0, j), h, 6, True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for i, (k, usd, uf, clp, fill) in enumerate(val_filas, start=1):
        for j, v in enumerate((k, usd, uf, clp)):
            if fill:
                shade(t5.cell(i, j), fill)
            celda(t5.cell(i, j), v, 6, j == 0,
                  align=WD_ALIGN_PARAGRAPH.RIGHT if j else None)

    parrafo(doc, DECLARACION, size=5.5)


def hoja2(doc):
    encabezado_hoja(doc, 2)
    barra(doc, "Mapa de ubicación de referencias", size=8, fill=GRIS, color=NEGRO)
    imagen(doc, "{d.mapa.staticMapUrl}", 18.0, 7.0)
    tf = doc.add_table(rows=2, cols=3)
    tf.autofit = False
    bordes(tf)
    for j in range(3):
        tf.columns[j].width = Cm(6.3)
        shade(tf.cell(0, j), GRIS)
        celda(tf.cell(0, j), "Fachada", 7, True, align=WD_ALIGN_PARAGRAPH.CENTER)
        imagen(tf.cell(1, j), None, 5.2, 4.0)  # sin campo foto por comparable (E-121..123)
    fichas = doc.add_table(rows=8, cols=6)
    fichas.autofit = False
    bordes(fichas)
    anchos_f = (1.6, 4.7, 1.6, 4.7, 1.6, 4.7)
    for fila in fichas.rows:
        for i, c in enumerate(fila.cells):
            c.width = Cm(anchos_f[i])
    for k in range(3):
        c0 = 2 * k
        fichas.cell(0, c0).merge(fichas.cell(0, c0 + 1))
        shade(fichas.cell(0, c0), GRIS)
        celda(fichas.cell(0, c0), f"Referencia Nº {k + 1}", 7, True,
              align=WD_ALIGN_PARAGRAPH.CENTER)
        idx = f"[i={k}]"
        pares = [
            ("Tipo:", f"{{d.comparablesInforme.filas{idx}.tipoReferencia}}"),
            ("Dirección:", f"{{d.comparablesInforme.filas{idx}.direccion}}"),
            ("Comuna:", f"{{d.comparablesInforme.filas{idx}.comuna}}"),
            ("Valor UF:", f"{{d.comparablesInforme.filas{idx}.precioUf:formatN(0)}}"),
            ("Año:", f"{{d.comparablesInforme.filas{idx}.anio}}"),
            ("Sup const:", f"{{d.comparablesInforme.filas{idx}.supConstruida:formatN(2)}}"),
            ("Terreno:", f"{{d.comparablesInforme.filas{idx}.supTerreno:formatN(2)}}"),
        ]
        for i, (kk, vv) in enumerate(pares, start=1):
            celda(fichas.cell(i, c0), kk, 6.5, True)
            celda(fichas.cell(i, c0 + 1), vv, 6.5)


def hoja3(doc):
    encabezado_hoja(doc, 3)

    def bloque(rotulo, pares, cols=4, tags=None):
        """pares = lista de rótulos fijos (valores vacíos salvo tags[rótulo])."""
        tags = tags or {}
        filas = (len(pares) + cols // 2 - 1) // (cols // 2)
        t = bloque_con_sidebar(doc, rotulo, filas, cols, None)
        idx = 0
        for f in range(filas):
            for cpar in range(cols // 2):
                if idx >= len(pares):
                    break
                k = pares[idx]
                cl = t.cell(f, cpar * 2)
                shade(cl, GRIS)
                celda(cl, k, 5.5, True)
                celda(t.cell(f, cpar * 2 + 1), tags.get(k, ""), 5.5, True)
                idx += 1
        return t

    bloque("Exigencias", [
        "Superf. Predial Mínimo", "Frente Predial Mínimo", "% Ocupación Suelo",
        "Sistema Agrupamiento", "Coef Máx Constructibidad", "Altura Edificación",
        "Distancia medianero", "Antejardin",
    ])
    bloque("Exigencias", [
        "D.F.L. 2", "LEY 6071 (Venta por Pisos)", "LEY 9135 (Ley Pereira)",
        "LEY 19537 (Coprop. Inmob.)", "Según el uso actual del bien",
        "Según tipo construcción", "Uso más probable del bien", "Cambios en Plan Regulador",
    ])
    bloque("Sector", [
        "Demanda", "Tendencia", "Densidad", "Mercado", "ABC1", "Destino barrio", "Tipo zona",
        "Calzada", "Acera", "Solera", "Antejardin", "Red Eléctrica", "Agua Potable", "Gas",
        "Proy. Uso Terrenos", "Casas", "Deptos.", "Condominios", "Equip. Comer", "Industrial",
        "Comercial", "Sitios Eriazos", "Otros",
    ])
    bloque("Arteria", ["Arteria Principal"], cols=2)
    bloque("Terreno", [
        "Superficie", "Forma", "Pendiente", "Orientación", "Frente", "Contrafrente",
        "NORTE", "SUR",
    ], tags={
        "Superficie": "{d.propiedad.supTerrenoM2:formatN(2)} m²",
        "Orientación": "{d.propiedad.orientacion}",
    })
    bloque("Emplazamiento", [
        "Tipo agrupamiento", "Diseño Arquitectónico", "Utilidad Funcional", "Tipo Adosamiento",
        "Calidad constructiva", "Estado de conservación", "Predio y/o vista",
        "Orientación construcción", "Relación Terr/Constr", "Iluminación natural",
    ], tags={
        "Tipo agrupamiento": "{d.propiedad.agrupacion}",
        "Calidad constructiva": "{d.propiedad.calidadConstruccion}",
        "Estado de conservación": "{d.propiedad.estadoConservacion:upperCase}",
    })

    # Características constructivas (4 col: Elemento | Materialidad | Calidad | Estado)
    elems = ["Estructura Soportante", "Divisiones Interiores", "Entrepisos", "Cubierta",
             "Revestimiento exterior", "Cierros exteriores", "O. Complementarias",
             "Construcción anexo", "Aire Acondicionado", "Calefacción", "Closet Mural",
             "Muebles de cocina", "Sanitarios", "Grifería", "Puerta Principal", "Ventanas"]
    cc = bloque_con_sidebar(doc, "Características Constructivas", len(elems) + 1, 4,
                            (4.5, 7.4, 3.5, 3.5))
    for j, h in enumerate(("Elementos", "Materialidad/Tipo", "Calidad", "Estado")):
        shade(cc.cell(0, j), GRIS)
        celda(cc.cell(0, j), h, 5.5, True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for i, e in enumerate(elems, start=1):
        celda(cc.cell(i, 0), e, 5.5, True)
        if e == "Estructura Soportante":
            celda(cc.cell(i, 1), "{d.propiedad.materialPredominante}", 5.5)

    # Terminaciones (loop)
    term = bloque_con_sidebar(doc, "Terminaciones", 4, 7,
                              (3.0, 3.0, 3.4, 2.9, 2.9, 1.9, 1.8))
    heads_t = ("Terminaciones", "Tipo de Pavimento", "Material/Marca/Origen",
               "Revestimiento de Muros", "Terminación de Cielo", "Iluminación", "Estado")
    for j, h in enumerate(heads_t):
        shade(term.cell(0, j), GRIS)
        celda(term.cell(0, j), h, 5.5, True, align=WD_ALIGN_PARAGRAPH.CENTER)
    celda(term.cell(1, 0), "{d.recintos.terminacionesPorRecinto[i].nombre}", 5.5)
    celda(term.cell(1, 1), "{d.recintos.terminacionesPorRecinto[i].categoria}", 5.5)
    celda(term.cell(1, 2), "{d.recintos.terminacionesPorRecinto[i].descripcion}", 5.5)
    celda(term.cell(1, 6), "{d.recintos.terminacionesPorRecinto[i].calidad}", 5.5)
    celda(term.cell(2, 0), "{d.recintos.terminacionesPorRecinto[i+1].nombre}", 5.5)
    for j, s in enumerate(("Alcantarillado", "Agua Potable", "Electricidad", "Gas", "Otros")):
        cl = term.cell(3, j)
        shade(cl, GRIS)
        celda(cl, s, 5.5, True)

    # Habitaciones (loop lista — la matriz 14×5 del oráculo no es reproducible con contrato plano)
    hab = bloque_con_sidebar(doc, "Habitaciones", 4, 3, (6.0, 8.0, 4.9))
    for j, h in enumerate(("Nivel", "Recinto", "Cantidad")):
        shade(hab.cell(0, j), GRIS)
        celda(hab.cell(0, j), h, 5.5, True, align=WD_ALIGN_PARAGRAPH.CENTER)
    celda(hab.cell(1, 0), "{d.recintos.habitacionesPorNivel[i].nivel}", 5.5)
    celda(hab.cell(1, 1), "{d.recintos.habitacionesPorNivel[i].tipoRecinto}", 5.5)
    celda(hab.cell(1, 2), "{d.recintos.habitacionesPorNivel[i].cantidad}", 5.5,
          align=WD_ALIGN_PARAGRAPH.RIGHT)
    celda(hab.cell(2, 0), "{d.recintos.habitacionesPorNivel[i+1].nivel}", 5.5)
    celda_partes(hab.cell(3, 0), [("Dormitorios: ", 5.5, True, None),
                                  ("{d.propiedad.dormitorios}", 5.5, True, None)])
    celda_partes(hab.cell(3, 1), [("Baños: ", 5.5, True, None),
                                  ("{d.propiedad.banos}", 5.5, True, None)])
    celda_partes(hab.cell(3, 2), [("Total: ", 5.5, True, None),
                                  ("{d.propiedad.supConstruccionM2:formatN(2)}", 5.5, True, None),
                                  (" m²", 5.5, True, None)])

    bloque("COMODIDADES", [
        "Muebles de Cocina", "Comedor de Diario", "Patio de servicio", "Estacionamiento",
        "Piscina", "Gimnasio", "Sauna", "Bodega", "Jardin conformado", "Calefacción", "Alarma",
        "Protecciones/Rejas", "Aspiración central", "Climatización", "Purificador de aire",
        "Corrientes. Débiles",
    ], cols=8)

    amp = bloque_con_sidebar(doc, "Ampliaciones", 3, 3, (8.0, 5.5, 5.4))
    for j, h in enumerate(("Descripción", "Metros Cuadrados", "Año regularización")):
        shade(amp.cell(0, j), GRIS)
        celda(amp.cell(0, j), h, 5.5, True, align=WD_ALIGN_PARAGRAPH.CENTER)
    celda(amp.cell(1, 0), "{d.recintos.ampliaciones[i].descripcion}", 5.5)
    celda(amp.cell(1, 1), "{d.recintos.ampliaciones[i].supM2:formatN(2)}", 5.5,
          align=WD_ALIGN_PARAGRAPH.RIGHT)
    celda(amp.cell(1, 2), "{d.recintos.ampliaciones[i].annoRegularizacion}", 5.5,
          align=WD_ALIGN_PARAGRAPH.RIGHT)
    celda(amp.cell(2, 0), "{d.recintos.ampliaciones[i+1].descripcion}", 5.5)


def hoja_fotos(doc, n):
    encabezado_hoja(doc, n)
    barra(doc, "Fotos de la Propiedad", size=8, fill=GRIS, color=NEGRO)
    if n == 4:
        # galería única con loop (la grilla 2×4 exacta no es reproducible con loop Carbone:
        # lista vertical foto+rótulo — desviación declarada)
        g = doc.add_table(rows=3, cols=1)
        g.autofit = False
        bordes(g)
        g.columns[0].width = Cm(19.5)
        imagen(g.cell(0, 0), "{d.fotos.fotos[i].url}", 9.1, 5.9)
        shade(g.cell(1, 0), GRIS)
        celda(g.cell(1, 0), "{d.fotos.fotos[i].categoria}", 7, True,
              align=WD_ALIGN_PARAGRAPH.CENTER)
        celda(g.cell(2, 0), "{d.fotos.fotos[i+1].categoria}", 7,
              align=WD_ALIGN_PARAGRAPH.CENTER)
    else:
        parrafo(doc, "(continuación de la galería de fotos)", size=7,
                align=WD_ALIGN_PARAGRAPH.CENTER)


def hoja_anexos(doc, n, numero):
    encabezado_hoja(doc, n)
    barra(doc, f"ANEXO N°{numero}", size=9)
    t0 = doc.add_table(rows=1, cols=2)
    t0.autofit = False
    bordes(t0)
    t0.columns[0].width = Cm(5.0)
    t0.columns[1].width = Cm(14.5)
    celda(t0.cell(0, 0), f"Anexo N° {numero}: Comentarios", 6.5, True)
    celda(t0.cell(0, 1), "", 6.5)
    rotulos = (["PLANO", "ESQUEMA DE SUPERFICIES", "INFORMACION DE SII"] if numero == 1 else
               ["ROL - AVALUO", "PERMISO DE EDIFICACIÓN", "NO EXPROPIACIÓN SERVIU",
                "RECEPCION FINAL"])
    for r in rotulos:
        barra(doc, r, size=7, fill=GRIS, color=NEGRO)
    if numero == 1:
        parrafo(doc, "Documentos del expediente (los archivos adjuntos se listan a continuación; "
                     "su inserción como imagen queda para el merge post-render):", size=6)
        ta = doc.add_table(rows=3, cols=2)
        ta.autofit = False
        bordes(ta)
        ta.columns[0].width = Cm(12.0)
        ta.columns[1].width = Cm(7.5)
        shade(ta.cell(0, 0), GRIS)
        shade(ta.cell(0, 1), GRIS)
        celda(ta.cell(0, 0), "Documento", 6.5, True)
        celda(ta.cell(0, 1), "Tipo", 6.5, True)
        celda(ta.cell(1, 0), "{d.anexos.documentos[i].nombre}", 6.5)
        celda(ta.cell(1, 1), "{d.anexos.documentos[i].tipo}", 6.5)
        celda(ta.cell(2, 0), "{d.anexos.documentos[i+1].nombre}", 6.5)


# ----------------------------------------------------------------------------- verificación


def normalizar(tag):
    t = re.sub(r"\[i(\+1)?(=\d+)?(,[^\]]*)?\]", "[i]", tag)
    t = re.sub(r":[^{}]*}$", "}", t)
    return t


def extraer_tags(docx_path):
    with zipfile.ZipFile(docx_path) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    descr = re.findall(r'descr="([^"]*)"', xml)
    texto = re.sub(r"<[^>]+>", "", xml)
    crudos = re.findall(r"\{d\.[^{}]*\}", texto)
    for d in descr:
        crudos += re.findall(r"\{d\.[^{}]*\}", d)
    return crudos


def verificar():
    requeridas = [ln.strip() for ln in REQUERIDAS.read_text(encoding="utf-8").splitlines()
                  if ln.strip()]
    presentes = {normalizar(t) for t in extraer_tags(SALIDA)}
    faltan = [r for r in requeridas if r not in presentes and r not in CUALITATIVA_EXCLUIDOS]
    print(f"verificar(): {len(requeridas) - len(faltan)}/{len(requeridas)} requeridas presentes; "
          f"{len(presentes)} tags únicos en el docx")
    if faltan:
        print("FALTAN:")
        for f in faltan:
            print("  " + f)
        sys.exit(1)


def _fijar_tabla(tbl, tope_dxa, W):
    """Fija layout fixed + tblW + tblGrid + tcW (lxml element), reescalando
    proporcionalmente si excede `tope_dxa`. Recurre en tablas anidadas.
    Defensa contra el autolayout de LibreOffice (render de Carbone), que
    infla columnas por contenido y RECORTA lo que cae fuera de la página."""
    filas = tbl.findall(f"{W}tr")
    if not filas:
        return

    def spans(tr):
        out = []
        for tc in tr.findall(f"{W}tc"):
            gs = tc.find(f"{W}tcPr/{W}gridSpan")
            out.append(int(gs.get(f"{W}val")) if gs is not None else 1)
        return out

    ncols = max(sum(spans(tr)) for tr in filas)

    # anchos: primera fila sin merges con tcW dxa completo; si no, gridCol; si no, iguales
    anchos = None
    for tr in filas:
        sp = spans(tr)
        if sum(sp) == ncols and len(sp) == ncols:
            ws = []
            for tc in tr.findall(f"{W}tc"):
                tcw = tc.find(f"{W}tcPr/{W}tcW")
                if tcw is None or tcw.get(f"{W}type") != "dxa":
                    ws = None
                    break
                ws.append(int(tcw.get(f"{W}w")))
            if ws:
                anchos = ws
                break
    if anchos is None:
        grid = tbl.find(f"{W}tblGrid")
        if grid is not None:
            gcs = [gc.get(f"{W}w") for gc in grid.findall(f"{W}gridCol")]
            if len(gcs) == ncols and all(gcs):
                anchos = [int(g) for g in gcs]
    if anchos is None:
        anchos = [tope_dxa // ncols] * ncols

    total = sum(anchos)
    if total > tope_dxa:
        anchos = [max(200, a * tope_dxa // total) for a in anchos]
        total = sum(anchos)

    tblPr = tbl.find(f"{W}tblPr")
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    for tag in ("w:tblW", "w:tblLayout"):
        for el in tblPr.findall(qn(tag)):
            tblPr.remove(el)
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), str(total))
    tblW.set(qn("w:type"), "dxa")
    tblPr.append(tblW)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)

    grid = tbl.find(f"{W}tblGrid")
    if grid is not None:
        tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for a in anchos:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(a))
        grid.append(gc)
    tbl.insert(list(tbl).index(tblPr) + 1, grid)

    for tr in filas:
        col = 0
        for tc in tr.findall(f"{W}tc"):
            gs = tc.find(f"{W}tcPr/{W}gridSpan")
            n = int(gs.get(f"{W}val")) if gs is not None else 1
            w = sum(anchos[col:col + n])
            col += n
            tcPr = tc.find(f"{W}tcPr")
            if tcPr is None:
                tcPr = OxmlElement("w:tcPr")
                tc.insert(0, tcPr)
            for el in tcPr.findall(qn("w:tcW")):
                tcPr.remove(el)
            tcW = OxmlElement("w:tcW")
            tcW.set(qn("w:w"), str(w))
            tcW.set(qn("w:type"), "dxa")
            tcPr.append(tcW)
            # anidadas: caben en el ancho del cell menos márgenes (~2×108 dxa)
            for sub in tc.findall(f"{W}tbl"):
                _fijar_tabla(sub, w - 250, W)


def enforce_fixed_layout(doc):
    """Todas las tablas top-level del body, con el tope de su sección:
    portada (antes del sectPr de párrafo) 18,8 cm; hojas 19,8 cm."""
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    tope = int(18.8 * 567)
    for child in doc.element.body:
        if child.tag == f"{W}p" and child.find(f"{W}pPr/{W}sectPr") is not None:
            tope = int(19.8 * 567)
        elif child.tag == f"{W}tbl":
            _fijar_tabla(child, tope, W)


def validar_anchos(path):
    """Validación aritmética post-guardado: ninguna tabla excede su tope."""
    from lxml import etree
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    root = etree.fromstring(zipfile.ZipFile(path).read("word/document.xml"))
    body = root.find(f"{W}body")
    tope = int(18.8 * 567)
    excesos = 0
    i = 0
    print("validar_anchos (tablas top-level):")
    for child in body:
        if child.tag == f"{W}p" and child.find(f"{W}pPr/{W}sectPr") is not None:
            tope = int(19.8 * 567)
        elif child.tag == f"{W}tbl":
            i += 1
            grid = child.find(f"{W}tblGrid")
            suma = sum(int(g.get(f"{W}w")) for g in grid.findall(f"{W}gridCol"))
            ncol = len(grid.findall(f"{W}gridCol"))
            ok = "OK" if suma <= tope + 5 else "EXCEDE"
            if ok == "EXCEDE":
                excesos += 1
            print(f"  tabla {i:2d}: cols={ncol:2d} suma={suma:5d} dxa (tope {tope}) {ok}")
    print(f"validar_anchos: {excesos} EXCEDEN de {i}")
    return excesos


def main():
    global PNG_GRIS
    PNG_GRIS = png_gris(AQUI / "_placeholder_gris.png")
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(6)
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.space_before = Pt(0)

    s0 = doc.sections[0]
    s0.page_width, s0.page_height = Cm(21), Cm(29.7)
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s0, m, Cm(1.0))
    portada(doc)

    s1 = doc.add_section(WD_SECTION.NEW_PAGE)
    s1.page_width, s1.page_height = Cm(21), Cm(29.7)
    s1.left_margin = s1.right_margin = Cm(0.5)
    s1.top_margin, s1.bottom_margin = Cm(0.5), Cm(0.4)

    hoja1(doc)
    doc.add_page_break()
    hoja2(doc)
    doc.add_page_break()
    hoja3(doc)
    doc.add_page_break()
    hoja_fotos(doc, 4)
    doc.add_page_break()
    hoja_fotos(doc, 5)
    doc.add_page_break()
    hoja_anexos(doc, 6, 1)
    doc.add_page_break()
    hoja_anexos(doc, 7, 2)

    enforce_fixed_layout(doc)
    doc.save(SALIDA)
    print(f"OK: {SALIDA} ({SALIDA.stat().st_size} bytes)")
    if validar_anchos(SALIDA) != 0:
        raise SystemExit("hay tablas que exceden el ancho útil")
    verificar()


if __name__ == "__main__":
    main()
