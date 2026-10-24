"""Celdas con fórmula de la hoja Impresion en un rango de filas: dirección, valor y fórmula."""
import sys, warnings
import openpyxl

warnings.filterwarnings("ignore")
ruta, desde, hasta = sys.argv[1], int(sys.argv[2]), int(sys.argv[3])
wf = openpyxl.load_workbook(ruta, keep_vba=True)["Impresion"]
wv = openpyxl.load_workbook(ruta, keep_vba=True, data_only=True)["Impresion"]
for fila in wf.iter_rows(min_row=desde, max_row=hasta):
    for c in fila:
        if isinstance(c.value, str) and c.value.startswith("="):
            v = wv[c.coordinate].value
            if v in (None, "", 0) and "IF(" in c.value.upper():
                estado = "vacía"
            else:
                estado = ""
            print(f"{c.coordinate:6} {str(v)[:34]!r:38} {c.value[:80]} {estado}")
