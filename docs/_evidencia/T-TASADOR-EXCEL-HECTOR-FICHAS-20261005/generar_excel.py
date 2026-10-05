# -*- coding: utf-8 -*-
"""
T-TASADOR-EXCEL-HECTOR-FICHAS-20261005 · Genera el Excel único para Héctor.
Fuente de todos los datos: docs/_evidencia/T-TASADOR-POBLAR-FICHAS-CLIENTE-DESDE-PLANILLA-20261005/
(mapeo.md · LISTA_FALTANTES.md · LISTA_DUPLICADOS.md · extraccion-planilla.json).
SOLO LECTURA de la evidencia; cero contacto con Airtable.
Salida: docs/_analisis/LISTA_HECTOR_Fichas_Clientes_20261005.xlsx
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

# ---------- Paleta (CLAUDE.md §4.4) ----------
AZUL = "075899"          # azul VProperty (encabezados)
AZUL_CLARO = "D9E7F3"    # secciones
NARANJA_RESP = "FDEBC8"  # celda de respuesta de Héctor (tinte del naranja de marca)
VERDE_REC = "E3F2E7"     # recomendación del equipo
GRIS_INFO = "F2F2F2"     # filas informativas
AMBAR_AVISO = "FFF4E0"   # avisos
ROJO_TXT = "B91C1C"

F_TITULO = Font(name="Calibri", size=16, bold=True, color=AZUL)
F_SUB = Font(name="Calibri", size=11, italic=True, color="444444")
F_HDR = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
F_SEC = Font(name="Calibri", size=12, bold=True, color=AZUL)
F_NORM = Font(name="Calibri", size=11)
F_BOLD = Font(name="Calibri", size=11, bold=True)
F_REC = Font(name="Calibri", size=11, bold=True, color="14532D")
F_AVISO = Font(name="Calibri", size=11, bold=True, color="92400E")

FILL_HDR = PatternFill("solid", fgColor=AZUL)
FILL_SEC = PatternFill("solid", fgColor=AZUL_CLARO)
FILL_RESP = PatternFill("solid", fgColor=NARANJA_RESP)
FILL_REC = PatternFill("solid", fgColor=VERDE_REC)
FILL_INFO = PatternFill("solid", fgColor=GRIS_INFO)
FILL_AVISO = PatternFill("solid", fgColor=AMBAR_AVISO)

THIN = Side(style="thin", color="B7C3CE")
BORDE = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
WRAP_C = Alignment(wrap_text=True, vertical="center", horizontal="center")

def fila(ws, r, valores, fills=None, fonts=None, borde=True):
    for c, v in enumerate(valores, start=1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.alignment = WRAP
        cell.font = F_NORM
        if borde:
            cell.border = BORDE
    if fills:
        for c, f in fills.items():
            ws.cell(row=r, column=c).fill = f
    if fonts:
        for c, f in fonts.items():
            ws.cell(row=r, column=c).font = f

def encabezado(ws, r, textos):
    for c, t in enumerate(textos, start=1):
        cell = ws.cell(row=r, column=c, value=t)
        cell.font = F_HDR
        cell.fill = FILL_HDR
        cell.alignment = WRAP_C
        cell.border = BORDE
    ws.row_dimensions[r].height = 30

def seccion(ws, r, texto, ncols):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
    cell = ws.cell(row=r, column=1, value=texto)
    cell.font = F_SEC
    cell.fill = FILL_SEC
    cell.alignment = WRAP
    for c in range(1, ncols + 1):
        ws.cell(row=r, column=c).border = BORDE
    ws.row_dimensions[r].height = 22

def preparar_impresion(ws, fila_hdr=None):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    if fila_hdr:
        ws.print_title_rows = f"{fila_hdr}:{fila_hdr}"

wb = Workbook()

# =====================================================================
# HOJA 1 · Confirmar
# =====================================================================
ws = wb.active
ws.title = "1. Confirmar"
ws.sheet_properties.tabColor = AZUL
anchos = [55, 18, 22, 45]
for i, a in enumerate(anchos, start=1):
    ws.column_dimensions[get_column_letter(i)].width = a

ws.merge_cells("A1:D1")
ws["A1"] = "LISTA PARA HÉCTOR — Fichas de clientes (factores y tasas) · VProperty"
ws["A1"].font = F_TITULO
ws.merge_cells("A2:D2")
ws["A2"] = ("Fecha: 05-10-2026 · Cómo usar este archivo: responde en las celdas naranjas de cada hoja. "
            "Con tus respuestas se destraba la carga de las fichas de clientes en el sistema.")
ws["A2"].font = F_SUB
ws["A2"].alignment = WRAP
ws.row_dimensions[2].height = 30

r = 4
seccion(ws, r, "LAS 2 DECISIONES QUE DESTRABAN TODO", 4)
r += 1
encabezado(ws, r, ["Pregunta", "Respuesta de Héctor", "Respuesta esperada", "Comentario de Héctor (opcional)"])
hdr1 = r
r += 1
fila(ws, r, [
    "1) ¿La fuente oficial de los factores y tasas de cada cliente es el formato de cálculo "
    "\"Formato-Informe-VProperty-Enero2026\" (el Excel con el que tasan todos los tasadores)? "
    "Si confirmas que SÍ, 20 clientes se cargan de inmediato con los valores de ese formato.",
    "", "SÍ / NO", ""],
    fills={2: FILL_RESP, 4: FILL_RESP}, fonts={1: F_NORM})
ws.row_dimensions[r].height = 68
r += 1
fila(ws, r, [
    "2) Regla del factor de seguro según tipo de propiedad: detectamos que el factor base del cliente "
    "sube a 1,0 cuando la propiedad es CASA (ejemplo: MetLife tiene base 0,8 y en una casa pasa a 1,0). "
    "¿Es correcta esta regla y aplica igual a todos los clientes?",
    "", "SÍ / NO · si varía, explicar cómo", ""],
    fills={2: FILL_RESP, 4: FILL_RESP})
ws.row_dimensions[r].height = 62
r += 2

seccion(ws, r, "PREGUNTA ADICIONAL (no bloquea la carga)", 4)
r += 1
fila(ws, r, [
    "3) Redondeo del valor final: ¿existe una regla de redondeo por cliente o ese dato se abandona? "
    "(hoy casi ninguna ficha lo tiene y no aparece en el formato de cálculo).",
    "", "Hay regla / Se abandona", ""],
    fills={2: FILL_RESP, 4: FILL_RESP})
ws.row_dimensions[r].height = 48
r += 2

ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
cell = ws.cell(row=r, column=1, value=(
    "AVISO (no es pregunta — solo para que lo sepas): 37 fichas tienen hoy un factor de seguro 0,825 que es un "
    "dato mal sembrado — el formato de cálculo nunca produce 0,825 (solo 0,8 o 1,0). Se corregirá "
    "automáticamente al cargar los valores correctos. No requiere ninguna acción tuya."))
cell.font = F_AVISO
cell.fill = FILL_AVISO
cell.alignment = WRAP
for c in range(1, 5):
    ws.cell(row=r, column=c).border = BORDE
ws.row_dimensions[r].height = 54
preparar_impresion(ws)

# =====================================================================
# HOJA 2 · Clientes a confirmar
# =====================================================================
ws = wb.create_sheet("2. Clientes a confirmar")
ws.sheet_properties.tabColor = "F5A213"
anchos = [5, 30, 26, 60, 24]
for i, a in enumerate(anchos, start=1):
    ws.column_dimensions[get_column_letter(i)].width = a

ws.merge_cells("A1:E1")
ws["A1"] = "Clientes que hoy NO se pueden cargar — una pregunta por fila"
ws["A1"].font = F_TITULO
ws.merge_cells("A2:E2")
ws["A2"] = "Responde en la columna naranja. Donde la pregunta ofrece opciones, basta con elegir una."
ws["A2"].font = F_SUB

HDR2 = ["Nº", "Cliente (ficha en el sistema)", "¿A qué cliente del formato corresponde?",
        "Pregunta para Héctor", "Respuesta de Héctor"]

r = 4
seccion(ws, r, "A · NOMBRES PARECIDOS — ¿ES EL MISMO CLIENTE? (17)", 5)
r += 1
encabezado(ws, r, HDR2)
hdr_row2 = r
r += 1

grupo_a = [
    ("\"ULH (BICE) Leasing\" y \"UNIDAD LEASING HABITACIONAL\"", "Unidad Leasing Habitacional",
     "¿\"Unidad Leasing Habitacional\" es lo mismo que \"ULH (BICE) Leasing\"? ¿Y qué relación tiene con las fichas BICE? (ver también la fila de Bice más abajo)"),
    ("\"M&V\" y \"M&V (Munita y Vergara)\"", "M&V",
     "Son dos fichas del mismo cliente: ¿cuál nombre queda como oficial?"),
    ("\"MAS LEASING\" y \"Másleasing\"", "MásLeasing",
     "Son dos fichas del mismo cliente: ¿cuál queda como oficial?"),
    ("\"ANDES\" y \"Administradora Andes S.A.\"", "Administradora Andes S.A.",
     "Son dos fichas del mismo cliente: ¿cuál queda como oficial?"),
    ("\"PARTICULARES\" y \"Particular\"", "Particular",
     "Son dos fichas del mismo cliente: ¿cuál queda como oficial? Ojo: \"Paris Crédito Hipotecario\" usa el mismo código abreviado (PAR) — ¿cómo los distinguimos?"),
    ("\"CHILE VIVIENDA\" y \"Chilevivienda\"", "Chile Vivienda",
     "Son dos fichas del mismo cliente: ¿cuál queda como oficial?"),
    ("\"CONCRECES\" y \"Concreces Leasing\"", "Concreces",
     "Son dos fichas del mismo cliente: ¿cuál queda como oficial? (solo \"Concreces Leasing\" trae la tasa correcta de 6,0%)"),
    ("\"BICE Hipotecaria\", \"BICE MUTUOS\" y \"BICE Leasing\"", "Bice Hipotecaria",
     "¿Cuál de las fichas BICE representa a \"Bice Hipotecaria\" del formato? ¿\"BICE Leasing\" y \"ULH (BICE) Leasing\" son líneas de negocio distintas con factores propios? Además, en el formato el código BCH está repetido entre Bice y Banco de Chile — ¿cuál es de quién?"),
    ("\"Banco de Chile\"", "Banco de Chile",
     "La ficha actual tiene factor de seguro 1,0 pero el formato dice base 0,8: ¿el 1,0 es intencional o es un error? (el código BCH repetido con Bice también afecta aquí)"),
    ("\"Coopeuch\"", "Copeuch",
     "El formato dice \"Copeuch\" y la ficha dice \"Coopeuch\": ¿es el mismo cliente (la cooperativa Coopeuch)?"),
    ("\"Penta Hipotecario\" y \"Penta Hipotecaria\"", "Penta Hipotecario",
     "Son dos fichas del mismo cliente: ¿cuál queda como oficial? (\"Penta Vida\" es un cliente DISTINTO y no se toca)"),
    ("\"NUEVO CAPITAL\" y \"Nuevo Capital Mutuos Hipotecarios\"", "Nuevo Capital Mutuos Hipotecarios",
     "Son dos fichas del mismo cliente: ¿cuál queda como oficial?"),
    ("\"CENTRAL HIPOTECARIA\" y \"Central Mutuos\"", "Central Mutuos",
     "¿\"CENTRAL HIPOTECARIA\" y \"Central Mutuos\" son el mismo cliente?"),
    ("\"TESSI\" y \"TESSI Servicios\"", "Tessi",
     "Son dos fichas del mismo cliente: ¿cuál queda como oficial? (solo \"TESSI Servicios\" trae la tasa de 6,0%)"),
    ("\"CREDITU\" y \"CrediTú\"", "CrediTú",
     "Son dos fichas del mismo cliente: ¿cuál queda como oficial?"),
    ("\"Santander Hipotecaria\"", "Banco Santander-Chile",
     "¿\"Santander Hipotecaria\" es el mismo cliente que \"Banco Santander-Chile\" del formato, o son dos clientes distintos?"),
    ("\"Scotiabank\" y \"Scotia Crédito Hipotecario\"", "Scotiabank",
     "Hay dos fichas Scotia: ¿cuál es la buena? ¿Son el mismo cliente?"),
]
n = 1
for ficha, formato, preg in grupo_a:
    fila(ws, r, [n, ficha, formato, preg, ""], fills={5: FILL_RESP})
    ws.row_dimensions[r].height = max(30, 16 * (len(preg) // 70 + 1))
    n += 1
    r += 1

r += 1
seccion(ws, r, "B · FICHAS SIN FUENTE EN NINGÚN LADO (17 fichas) — ¿cliente real y activo? ¿se carga con qué factores o se desactiva?", 5)
r += 1
encabezado(ws, r, HDR2)
r += 1
grupo_b = [
    ("Aseguradora Continental", "— (no aparece en el formato)",
     "¿Es cliente real de tasaciones? Su ficha dice seguro 1,0: ¿de dónde salió ese valor?"),
    ("BBVA Hipotecaria", "— (no aparece en el formato)",
     "¿Sigue vigente? (BBVA Chile no existe como banco desde 2018). ¿Se desactiva?"),
    ("Banco Falabella", "— (no aparece en el formato)",
     "¿Cliente vigente? ¿Con qué factores se carga o se desactiva?"),
    ("Banco Itaú", "— (no aparece en el formato)",
     "¿Cliente vigente? Su ficha dice seguro 1,0: ¿es real?"),
    ("Banco Ripley", "— (no aparece en el formato)",
     "¿Cliente vigente? ¿Con qué factores se carga o se desactiva?"),
    ("Citi Mutuos", "— (no aparece en el formato)",
     "¿Cliente vigente? ¿Con qué factores se carga o se desactiva?"),
    ("Eurocapital", "— (no aparece en el formato)",
     "¿Cliente vigente? ¿Con qué factores se carga o se desactiva?"),
    ("HCS / Hipotecaria Compass", "— (no aparece en el formato)",
     "¿Cliente vigente? ¿Con qué factores se carga o se desactiva?"),
    ("HLC Hipotecaria Continental", "— (no aparece en el formato)",
     "¿Qué relación tiene con \"Hipotecaria La Construcción\"? (hay tres nombres parecidos — ver sección C)"),
    ("Mi Hipoteca", "— (no aparece en el formato)",
     "¿Cliente vigente? ¿Con qué factores se carga o se desactiva?"),
    ("Mutuosa", "— (no aparece en el formato)",
     "¿Cliente vigente? ¿Con qué factores se carga o se desactiva?"),
    ("Paris Crédito Hipotecario", "— (no aparece en el formato)",
     "¿Cliente vigente? Su código abreviado choca con el de \"Particulares\" — ¿cómo lo distinguimos?"),
    ("Sim Hipotecaria", "— (no aparece en el formato)",
     "¿Cliente vigente? ¿Con qué factores se carga o se desactiva?"),
    ("VALÓN Hipotecaria (existe en DOS copias idénticas)", "— (no aparece en el formato)",
     "¿Cliente vigente? ¿Con qué factores se carga o se desactiva? (la copia sobrante se limpia en la hoja 3)"),
    ("Vivienda Plus", "— (no aparece en el formato)",
     "¿Cliente vigente? ¿Con qué factores se carga o se desactiva?"),
    ("Scotia Crédito Hipotecario", "¿Scotiabank?",
     "Ya preguntado en la sección A (fila 17): si es lo mismo que \"Scotiabank\", aquí no hay nada más que responder."),
]
for ficha, formato, preg in grupo_b:
    fila(ws, r, [n, ficha, formato, preg, ""], fills={5: FILL_RESP})
    ws.row_dimensions[r].height = 32
    n += 1
    r += 1

r += 1
seccion(ws, r, "C · APARECEN EN LA PLANILLA OPERATIVA PERO NO EN EL FORMATO DE CÁLCULO (3)", 5)
r += 1
encabezado(ws, r, HDR2)
r += 1
grupo_c = [
    ("HIPOTECARIA LA CONSTRUCCION", "— (sin fila en el formato)",
     "¿Sigue operativa? ¿Con qué nombre se tasa hoy y qué factor/tasa le corresponde? ¿Su ficha buena es \"HIPOTECARIA LA CONSTRUCCION\" o \"La Construcción Hipotecaria\"?"),
    ("EXTERIOR", "— (sin fila en el formato)",
     "¿Qué cliente es \"EXTERIOR\"? ¿Está activo? ¿Qué factores le corresponden?"),
    ("FINANCIERA Y HABITACIONAL", "— (sin fila en el formato)",
     "¿Está activo? ¿Bajo qué nombre del formato se tasa y con qué factores?"),
]
for ficha, formato, preg in grupo_c:
    fila(ws, r, [n, ficha, formato, preg, ""], fills={5: FILL_RESP})
    ws.row_dimensions[r].height = 40
    n += 1
    r += 1

r += 1
seccion(ws, r, "D · CONFLICTOS DE VALOR (2) — aquí hay dos fuentes que se contradicen y hay que elegir", 5)
r += 1
encabezado(ws, r, HDR2)
r += 1
grupo_d = [
    ("ICGE / ICGE Internacional", "ICGE",
     "La fórmula vigente del formato le da tasa 4,5%, pero la intención aparente (y la ficha \"ICGE Internacional\") dicen 6,0%. ¿Cuál es la tasa correcta de ICGE? Si es 6,0%, el formato también está malo y habría que corregirlo. ¿Y son ICGE e ICGE Internacional el mismo cliente?"),
    ("Leasing Urbano", "Leasing Urbano",
     "La ficha vigente tiene tasa 6,0% pero el formato NO lo incluye en la lista de 6,0% (le daría 4,5%). ¿Cuál es la correcta: 6,0% de la ficha o 4,5% del formato?"),
]
for ficha, formato, preg in grupo_d:
    fila(ws, r, [n, ficha, formato, preg, ""], fills={5: FILL_RESP})
    ws.row_dimensions[r].height = 56
    n += 1
    r += 1

ws.freeze_panes = f"A{hdr_row2 + 1}"
preparar_impresion(ws, hdr_row2)

# =====================================================================
# HOJA 3 · Duplicados a limpiar
# =====================================================================
ws = wb.create_sheet("3. Duplicados a limpiar")
ws.sheet_properties.tabColor = "B91C1C"
anchos = [5, 22, 48, 44, 14, 22]
for i, a in enumerate(anchos, start=1):
    ws.column_dimensions[get_column_letter(i)].width = a

ws.merge_cells("A1:F1")
ws["A1"] = "Fichas repetidas — qué recomienda conservar el equipo y tu visto bueno"
ws["A1"].font = F_TITULO
ws.merge_cells("A2:F2")
ws["A2"] = ("Nada se ha borrado todavía. Marca SÍ en \"¿OK Héctor?\" si estás de acuerdo con conservar la ficha "
            "recomendada (en verde) y limpiar las demás; si no, anota cuál prefieres en el comentario.")
ws["A2"].font = F_SUB
ws["A2"].alignment = WRAP
ws.row_dimensions[2].height = 30

HDR3 = ["Nº", "Cliente", "Fichas repetidas (cómo distinguirlas)",
        "Recomendación del equipo: CONSERVAR", "¿OK Héctor? (SÍ/NO)", "Comentario de Héctor"]

r = 4
seccion(ws, r, "A · CON FICHA EN USO CLARA (6 grupos) — recomendación firme", 6)
r += 1
encabezado(ws, r, HDR3)
hdr_row3 = r
r += 1

grupos_a = [
    ("MetLife",
     "3 fichas: \"MetLife\" (en uso, con 3 tasaciones vigentes) · \"METLIFE\" (vacía) · \"MetLife Chile S.A.\" (semilla; es la única ficha de toda la base con redondeo definido = 2, dato a rescatar antes de limpiar)",
     "\"MetLife\" (la que está en uso)"),
    ("Agencia Habitacional",
     "\"Agencia Habitacional\" (en uso, tasación VP-0074) · \"AGENCIA\" (solo una solicitud cancelada)",
     "\"Agencia Habitacional\""),
    ("Austral Leasing",
     "\"Austral Leasing Habitacional\" (en uso, tasación VP-0075) · \"LEASING AUSTRAL\" (vacía)",
     "\"Austral Leasing Habitacional\""),
    ("Hipotecaria Security",
     "\"Hipotecaria Security S.A.\" (en uso, tasación VP-0076; le falta el código abreviado) · \"Hipotecaria Security\" (sin uso) · \"SECURITY PRINCIPAL\" (sin uso) — todo indica que las tres son el mismo grupo; confirmar que \"SECURITY PRINCIPAL\" no es un cliente aparte",
     "\"Hipotecaria Security S.A.\" (la que está en uso; completarle el código)"),
    ("Hipotecaria Evoluciona",
     "\"Hipotecaria Evoluciona\" (en uso, tasación VP-0077) · \"EVOLUCIONA\" (6 solicitudes, TODAS canceladas — el conteo engaña)",
     "\"Hipotecaria Evoluciona\""),
    ("4 Life",
     "\"4 LIFE\" (en uso, tasación VP-0062 visitada) · \"4LIFE\" (vacía)",
     "\"4 LIFE\""),
]
n = 1
for cli, fichas, rec in grupos_a:
    fila(ws, r, [n, cli, fichas, rec, "", ""],
         fills={4: FILL_REC, 5: FILL_RESP, 6: FILL_RESP}, fonts={4: F_REC})
    ws.row_dimensions[r].height = max(34, 15 * (len(fichas) // 55 + 1))
    n += 1
    r += 1

r += 1
seccion(ws, r, "B · PARES SIN TASACIONES VIGENTES (19 grupos) — recomendación tentativa, confirma cuál queda", 6)
r += 1
encabezado(ws, r, HDR3)
r += 1
grupos_b = [
    ("Afianza", "\"Afianza\" · \"AFIANZA\" (vacía, 2 solicitudes canceladas)", "\"Afianza\""),
    ("Banco de Chile", "\"Banco de Chile\" (con factores, pero seguro 1,0 en duda — ver hoja 2) · \"BANCO DE CHILE\" (vacía, con 10 solicitudes históricas canceladas)", "\"Banco de Chile\" (decidir antes el seguro 1,0 de la hoja 2)"),
    ("Valor Presente", "\"Valor Presente (VP)\" · \"VALOR PRESENTE (VLP)\" (vacía)", "\"Valor Presente (VP)\""),
    ("ServiHabit", "\"ServiHabit (SVH)\" · \"SERVIHABIT\" (vacía)", "\"ServiHabit (SVH)\""),
    ("Credihome", "\"Credihome (CRD)\" · \"CREDIHOME\" (vacía)", "\"Credihome (CRD)\""),
    ("VALÓN Hipotecaria", "Dos copias 100% idénticas (mismo nombre y mismos valores); solo cambia la fecha de creación", "La más antigua; eliminar la copia más nueva"),
    ("Administradora Andes", "\"Administradora Andes S.A.\" (1 solicitud cancelada) · \"ANDES\" (vacía)", "\"Administradora Andes S.A.\""),
    ("CrediTú", "\"CrediTú\" · \"CREDITU\" (valores idénticos en ambas)", "\"CrediTú\""),
    ("Chilevivienda", "\"Chilevivienda\" · \"CHILE VIVIENDA\" (vacía)", "\"Chilevivienda\""),
    ("Concreces", "\"Concreces Leasing\" (trae la tasa correcta de 6,0%) · \"CONCRECES\" (tasa 4,5%)", "\"Concreces Leasing\""),
    ("ICGE", "\"ICGE Internacional\" (tasa 6,0%) · \"ICGE\"", "A DECIDIR — depende del conflicto de tasa (hoja 2, sección D)"),
    ("La Construcción", "\"HIPOTECARIA LA CONSTRUCCION\" · \"La Construcción Hipotecaria\" · ver además \"HLC Hipotecaria Continental\" (¿tercer nombre del mismo cliente?)", "A DECIDIR — responde primero la hoja 2, sección C"),
    ("M&V", "\"M&V (Munita y Vergara)\" · \"M&V\"", "\"M&V (Munita y Vergara)\""),
    ("Más Leasing", "\"Másleasing\" (trae la tasa correcta de 6,0%) · \"MAS LEASING\" (tasa 4,5%)", "\"Másleasing\""),
    ("Nuevo Capital", "\"Nuevo Capital Mutuos Hipotecarios\" · \"NUEVO CAPITAL\" (vacía)", "\"Nuevo Capital Mutuos Hipotecarios\""),
    ("Particulares", "\"Particular\" · \"PARTICULARES\" (vacía) — ojo código compartido con \"Paris Crédito Hipotecario\"", "\"Particular\""),
    ("Penta", "\"Penta Hipotecario\" · \"Penta Hipotecaria\" (\"Penta Vida\" es OTRO cliente: no se toca)", "A DECIDIR — ¿cuál nombre es el correcto?"),
    ("TESSI", "\"TESSI Servicios\" (trae la tasa correcta de 6,0%) · \"TESSI\" (tasa 4,5%)", "\"TESSI Servicios\""),
    ("ULH / BICE", "Hasta 5 fichas para 2 clientes del formato: \"UNIDAD LEASING HABITACIONAL\" (vacía) · \"ULH (BICE) Leasing\" (valores correctos del formato) · \"BICE Hipotecaria\" · \"BICE Leasing\" · \"BICE MUTUOS\"", "A DECIDIR — el nudo más enredado; requiere tu respuesta en la hoja 2 (filas 1 y 8)"),
]
for cli, fichas, rec in grupos_b:
    tent = rec.startswith("A DECIDIR")
    fila(ws, r, [n, cli, fichas, rec, "", ""],
         fills={4: (FILL_AVISO if tent else FILL_REC), 5: FILL_RESP, 6: FILL_RESP},
         fonts={4: (F_AVISO if tent else F_REC)})
    ws.row_dimensions[r].height = max(30, 15 * (len(fichas) // 55 + 1))
    n += 1
    r += 1

r += 1
seccion(ws, r, "C · BASURA Y FICHAS DE PRUEBA (3)", 6)
r += 1
encabezado(ws, r, HDR3)
r += 1
grupos_c = [
    ("(fila basura)",
     "Una ficha cuyo \"nombre\" es literalmente la palabra \"nombre\": es la cabecera de un archivo que se importó por error como si fuera un cliente. No tiene ningún uso.",
     "ELIMINAR"),
    ("Fichas de prueba (2)",
     "Dos fichas \"SANDBOX\" creadas para pruebas técnicas en septiembre.",
     "Las maneja el equipo técnico; jamás se les cargan factores. No requieren tu respuesta."),
]
for cli, fichas, rec in grupos_c:
    fila(ws, r, [n, cli, fichas, rec, "", ""],
         fills={4: FILL_REC, 5: FILL_RESP, 6: FILL_RESP}, fonts={4: F_REC})
    ws.row_dimensions[r].height = 44
    n += 1
    r += 1

r += 1
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
cell = ws.cell(row=r, column=1, value=(
    "Nota del equipo: la limpieza se hará en una tanda aparte, con respaldo. Antes de eliminar fichas que tengan "
    "historial (aunque sea de solicitudes canceladas) se decidirá qué pasa con ese historial."))
cell.font = F_AVISO
cell.fill = FILL_AVISO
cell.alignment = WRAP
for c in range(1, 7):
    ws.cell(row=r, column=c).border = BORDE
ws.row_dimensions[r].height = 36
ws.freeze_panes = f"A{hdr_row3 + 1}"
preparar_impresion(ws, hdr_row3)

# =====================================================================
# HOJA 4 · Listos (informativo)
# =====================================================================
ws = wb.create_sheet("4. Listos (informativo)")
ws.sheet_properties.tabColor = "15803D"
anchos = [5, 34, 20, 18, 10, 48]
for i, a in enumerate(anchos, start=1):
    ws.column_dimensions[get_column_letter(i)].width = a

ws.merge_cells("A1:F1")
ws["A1"] = "SOLO INFORMATIVO — 20 clientes listos para cargar (no requieren acción tuya)"
ws["A1"].font = F_TITULO
ws.merge_cells("A2:F2")
ws["A2"] = ("Estos 20 clientes quedan cargados de inmediato con los valores del formato de cálculo apenas "
            "confirmes la pregunta 1 de la hoja \"1. Confirmar\". Se muestran para que veas el avance.")
ws["A2"].font = F_SUB
ws["A2"].alignment = WRAP
ws.row_dimensions[2].height = 30

r = 4
encabezado(ws, r, ["Nº", "Cliente", "Factor de garantía y seguro (base)",
                   "¿Sube a 1,0 si es casa?", "Tasa", "Nota"])
hdr_row4 = r
r += 1
listos = [
    ("Hipotecaria Security S.A.", "0,8", "Sí", "4,5%",
     "En un caso real se usó una tasa de 5,5% puesta a mano; se respeta el caso a caso."),
    ("Hipotecaria Evoluciona", "0,8", "Sí", "4,5%", ""),
    ("Agencia Habitacional", "1,0", "— (ya es 1,0)", "6,0%", "Valor 1,0 ya validado en producción."),
    ("Afianza", "1,0", "— (ya es 1,0)", "4,5%", ""),
    ("4 LIFE", "0,8", "Sí", "4,5%", ""),
    ("CrediHome", "1,0", "— (ya es 1,0)", "6,0%", ""),
    ("ServiHabit", "0,8", "Sí", "4,5%", ""),
    ("Credicasa Del Maule SPA", "1,0", "— (ya es 1,0)", "4,5%", ""),
    ("Austral Leasing Habitacional", "1,0", "— (ya es 1,0)", "4,5%", "Valor 1,0 ya validado en producción."),
    ("Tu Hipotecaria Chile", "1,0", "— (ya es 1,0)", "4,5%", ""),
    ("Casa Pronta", "1,0", "— (ya es 1,0)", "4,5%", ""),
    ("Ohio National", "0,8", "Sí", "4,5%", ""),
    ("Banco Estado", "0,8", "Sí", "4,5%", ""),
    ("Su Casa Hoy", "1,0", "— (ya es 1,0)", "4,5%", ""),
    ("Casa Nuestra", "1,0", "— (ya es 1,0)", "4,5%", ""),
    ("Penta Vida", "0,8", "Sí", "4,5%", "Cliente distinto de \"Penta Hipotecario\"."),
    ("Valor Presente", "0,8", "Sí", "4,5%", ""),
    ("Espacio Nuestro", "0,8", "Sí", "4,5%", ""),
    ("BCI", "0,8", "Sí", "4,5%", "La ficha del sistema se llama \"BCI Mutuos\"; mismo código BCI."),
    ("Consorcio", "0,8", "Sí", "4,5%", ""),
]
n = 1
for cli, base, casa, tasa, nota in listos:
    fila(ws, r, [n, cli, base, casa, tasa, nota], fills={c: FILL_INFO for c in range(1, 7)})
    ws.row_dimensions[r].height = 24
    n += 1
    r += 1

r += 1
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
cell = ws.cell(row=r, column=1, value=(
    "Recordatorio: \"Sube a 1,0 si es casa\" es la regla de la pregunta 2 de la hoja \"1. Confirmar\". "
    "Esa subida la aplica el sistema al calcular cada tasación; en la ficha se guarda el valor base."))
cell.font = F_AVISO
cell.fill = FILL_AVISO
cell.alignment = WRAP
for c in range(1, 7):
    ws.cell(row=r, column=c).border = BORDE
ws.row_dimensions[r].height = 36
ws.freeze_panes = f"A{hdr_row4 + 1}"
preparar_impresion(ws, hdr_row4)

salida = "/mnt/c/Users/Sergio/Documents/GitHub/if-ejecutiva/docs/_analisis/LISTA_HECTOR_Fichas_Clientes_20261005.xlsx"
wb.save(salida)
print("OK:", salida)
